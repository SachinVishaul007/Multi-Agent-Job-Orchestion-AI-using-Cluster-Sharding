#!/usr/bin/env python3
try:
    import openpyxl
    
    # Create a new workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    
    # Add headers
    ws['A1'] = 'Resume Points'
    ws['B1'] = 'Detected Tags'
    
    # Add sample data
    ws['A2'] = 'Developed RESTful APIs using Spring Boot and Java'
    ws['B2'] = 'java,spring boot,rest api'
    
    ws['A3'] = 'Implemented microservices architecture with Docker'
    ws['B3'] = 'docker,microservices,architecture'
    
    ws['A4'] = 'Built machine learning models using Python and TensorFlow'
    ws['B4'] = 'python,tensorflow,machine learning'
    
    # Save the file
    wb.save('test_resume.xlsx')
    print("Created test_resume.xlsx successfully")
    
except ImportError:
    print("openpyxl not available, creating a basic text file instead")
    with open('test_resume.txt', 'w') as f:
        f.write('Resume Points\tDetected Tags\n')
        f.write('Developed RESTful APIs using Spring Boot and Java\tjava,spring boot,rest api\n')
        f.write('Implemented microservices architecture with Docker\tdocker,microservices,architecture\n')
        f.write('Built machine learning models using Python and TensorFlow\tpython,tensorflow,machine learning\n')
    print("Created test_resume.txt instead")